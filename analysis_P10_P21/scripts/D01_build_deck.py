"""Collaborator deck: P10 BLA cell types, P10 vs P21, TFs, and cue/receptor
candidates for the Mef2c-cKO ectopic cortex -> BLA innervation question."""
import os
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT = Path(os.environ.get("BLA_ROOT", r"K:\scRNA_BLA_phd\P10_P21_amygdala_2026-10"))
F = ROOT / "figures"
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
DARK = RGBColor(0x1F, 0x2A, 0x44); GREY = RGBColor(0x55, 0x5B, 0x66); RED = RGBColor(0xB2, 0x18, 0x2B)


def text(slide, x, y, w, h, lines, size=14, color=DARK, bold_first=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        bullet = ln.startswith("- ")
        r = p.add_run(); r.text = ("\u2022 " + ln[2:]) if bullet else ln
        r.font.size = Pt(size); r.font.color.rgb = color
        r.font.bold = bold_first and i == 0
        p.space_after = Pt(5)
    return tb


def slide(title, sub=None):
    s = prs.slides.add_slide(BLANK)
    text(s, 0.45, 0.25, 12.4, 0.6, [title], size=24, bold_first=True)
    if sub: text(s, 0.45, 0.85, 12.4, 0.45, [sub], size=13, color=GREY)
    return s


def pic(s, name, x, y, w, h):
    """Place image inside the (x, y, w, h) box, keeping aspect ratio, centred."""
    p = F / name; W, H = Image.open(p).size
    sc = min(w / (W / 100), h / (H / 100)); ww, hh = W / 100 * sc, H / 100 * sc
    s.shapes.add_picture(str(p), Inches(x + (w - ww) / 2), Inches(y + (h - hh) / 2), Inches(ww), Inches(hh))


# 1 title
s = prs.slides.add_slide(BLANK)
text(s, 0.8, 2.2, 11.8, 1.2, ["Basolateral amygdala at P10: cell types, P10 vs P21, and candidate cues for Mef2c-cKO ectopic cortical axons"], size=30, bold_first=True)
text(s, 0.8, 3.9, 11.8, 1.6, ["Hansol Lim - BLA 10x scRNA-seq (P10 n=1, P21 n=2), 13,625 cells after QC",
                              "Reference: Allen Whole Mouse Brain atlas (Yao et al. 2023), amygdala ROIs CTXsp + sAMY (199k adult cells)",
                              "October 2026 - first-pass analysis for discussion"], size=15, color=GREY)

# 2 question + data
s = slide("The question and what this data can answer")
text(s, 0.5, 1.4, 6.1, 5.6, [
    "Collaborator's finding",
    "- Mef2c cKO: L2/3 callosal axons lose contralateral S1/S2 targeting",
    "- Ectopic axons appear around BLA / FoxP2+ ITCs from ~P5, increasing and branching at P8-P10",
    "- Source neurons may not be in SSc - still open",
    "Strategy: amygdala first, then back to cortex",
    "- Which BLA/ITC cell populations exist at P10?",
    "- Which of them express guidance / adhesion cues during P5-P10?",
    "- Which cortical neurons carry the matching axon-side partners?",
    "- Where is Mef2c itself expressed in the amygdala?"], size=14, bold_first=True)
text(s, 6.9, 1.4, 6.0, 5.6, [
    "Data",
    "- BLA dissections, 10x 3' v3 (Klein lab, processed by T. Straub, mm10-2020-A)",
    "- P10s1: 6,037 cells | P21s1: 2,507 | P21s2: 5,081 (after QC)",
    "- 3,079 cells removed: doublets (two-lineage co-expression) and low-quality clusters, logged per sample",
    "- Labels from marker genes plus correlation mapping to Allen subclasses / supertypes",
    "Caveats",
    "- P10 is ONE animal; P10 vs P21 differences include dissection and batch",
    "- The dissection includes CeA / MeA / BMA neighbours (expected for a BLA punch)",
    "- Allen reference and the cortex panel are ADULT; P10 neurons are immature"], size=14, bold_first=True)

# 3 UMAP
s = slide("35 cell types in P10 amygdala", "BLA glutamatergic (1-10), ITC / CeA / MeA GABA (11-16), interneurons (17-25), glia and vasculature (26-35)")
pic(s, "F1_P10_umap_cell_types.png", 0.3, 1.35, 12.7, 6.0)

# 4 markers
s = slide("Labels agree with known amygdala markers", "Dot size = % cells, colour = scaled mean expression, P10 cells only")
pic(s, "F2_P10_marker_dotplot.png", 0.2, 1.3, 9.6, 6.1)
text(s, 9.9, 1.4, 3.3, 6.0, [
    "Allen mapping (cells agreeing)",
    "- Rspo2/Etv1 BLA Glut -> 014 LA-BLA-BMA-PA Glut (92%)",
    "- ITC (a) / (b) -> 064 STR-PAL Chst9 Gaba (91% / 80%); Foxp2+ Tshz1+ Pbx3+ Meis2+",
    "- D2 SPN-like -> 062 STR D2 Gaba (83%)",
    "- Vip -> 046 Vip Gaba (93%); Pvalb chandelier -> 051 (72%)",
    "- OPC 99%, Microglia 85%, Ependymal 87%",
    "Other immature P10 glut types partly map to adult DG / IMN types; their labels rest on markers (next slide)"], size=11.5, bold_first=True)

# 5 glut subtypes
s = slide("BLA glutamatergic sub-populations", "Re-clustered glutamatergic neurons only (P10 + P21, Harmony across samples)")
pic(s, "F3_BLA_glut_subtypes.png", 0.2, 1.3, 8.4, 6.1)
text(s, 8.7, 1.4, 4.4, 6.0, [
    "Ten glutamatergic groups",
    "- Rspo2/Etv1: anterior-BLA Rspo2 population (Kim et al. 2016); also Cdh9, Sema3e, C1ql3",
    "- Tshz2/Satb1, Cdh8/Tshz3, Fgf10/Nrp1, Otof/Trhr, Vgll3/Prr16: Rspo2-negative BLA/LA groups",
    "- BMA Slc17a6/Zfp804b and PA/BMAp Esr1/Reln: neighbouring nuclei",
    "- Immature Igfbpl1/Epha3 (Dcx/Sox11+, still present at P21)",
    "- CLA/EPd-like Satb2/Nr4a2",
    "Ppp1r1b (posterior BLA) and Fezf2 are barely detected at P10/P21 in this dissection",
    "All groups are found at both ages"], size=12, bold_first=True)
s = slide("BLA glutamatergic marker genes on the glutamatergic map")
pic(s, "F3b_BLA_glut_featureplots.png", 2.6, 1.0, 8.2, 6.4)

# 6 composition
s = slide("Cell-type composition, P10 vs P21", "Dots = individual samples. With one P10 animal, read these as observations, not tests.")
pic(s, "F4_composition_P10_vs_P21.png", 0.2, 1.35, 12.9, 4.9)
text(s, 0.5, 6.3, 12.3, 1.1, [
    "- Neuronal share falls from P10 to P21 (more glia are captured as the brain matures)",
    "- CeA / D1 / D2 SPN-like fractions are higher at P10, most likely a wider medial punch; ITC fractions are higher at P21"], size=13)

# 7 DE
s = slide("Gene-expression changes P10 -> P21", "Wilcoxon per cell type, kept only if the direction holds against BOTH P21 replicates (pseudo-bulk)")
pic(s, "F5_DE_P10_vs_P21.png", 0.2, 1.35, 12.9, 5.4)
text(s, 0.5, 6.75, 12.3, 0.7, ["Most change is in BLA glutamatergic neurons. Shared P10-high genes include Ncam1, Sema6d, Cadm2, Grip1 and Pcdh7 (adhesion/guidance). P21-high genes include activity/maturation genes (Camk2a, Rasgrf1)."], size=12)

# 8 Mef2
s = slide("Mef2c in the amygdala", "Mef2c is high in most BLA glutamatergic groups, Pvalb chandelier cells and microglia; moderate in ITCs")
pic(s, "F6a_Mef2_family_dotplot.png", 0.2, 1.3, 6.4, 6.1)
pic(s, "F8_Mef2c.png", 6.7, 1.3, 6.4, 6.1)

# 9 TFs
s = slide("Transcription factors that define each P10 cell type", "Top 3 TFs per cell type from 1,321 curated mouse TFs (AnimalTFDB-based list from the GPCR/TF panel project)")
pic(s, "F6b_celltype_TFs_P10.png", 0.2, 1.3, 12.9, 6.1)

# 10 cues
s = slide("Guidance / adhesion cues expressed by amygdala neurons at P10")
pic(s, "F7a_guidance_cues_P10.png", 0.2, 1.0, 12.9, 6.3)
s = slide("Candidate cues that mark ITC and BLA sub-populations at P10", "Circled = >=30% of cells and >=1.6x / +10 pp over the mean of the other amygdala neuron types")
pic(s, "F7c_candidate_cues_ITC_BLAglut.png", 0.2, 1.35, 12.9, 4.9)
text(s, 0.5, 6.3, 12.3, 1.1, [
    "- ITC: Kirrel3 (98%), Sdk2, Sema6d, Sema5b, Cdh18, Ntng1 (both ITC groups) - homophilic adhesion and semaphorin cues",
    "- Rspo2/Etv1 BLA: C1ql3 (x10.7), Cdh8, Cdh9, Slit2, Sema3e. Other BLA groups: Cntn4/6, Ntng1/2, Slit3, Cdh13"], size=13)

# 11 cortex
s = slide("Back to cortex: which cortical neurons carry the matching partners?", "Adult Allen Isocortex, all areas pooled - a ranking aid only, not P5-P10 and not area-specific")
pic(s, "F9_cortex_receptors_adult_Allen.png", 0.2, 1.35, 12.9, 4.9)
text(s, 0.5, 6.3, 12.3, 1.1, [
    "- Kirrel3 (homophilic, ITC 98%): 90% of L2/3 IT and 99% of L5 ET cells, but only 37-39% of L5/L6 IT",
    "- Plxnd1 (receptor for Sema3e, made by Rspo2 BLA cells) is restricted to L4/5 IT (50%) and L2/3 IT (25%)",
    "- Many partners (Robo1/2, Lrrc4c, Adgrb3, Nrxn2) are near-ubiquitous in adult cortex, so they cannot single out a source population"], size=12.5)


# 12a scenarios overview
s = slide("Why would Mef2c-cKO axons over-innervate BLA? Six testable scenarios", "Hypotheses for discussion; slides that follow test scenarios 2 and 3 with this data")
text(s, 0.5, 1.4, 6.1, 5.8, [
    "Axon / cortex side",
    "- 1 Identity shift: Mef2c loss pushes L2/3-L4 IT neurons toward an amygdala-projecting IT programme (cf. Satb2 KO re-routing callosal axons; Alcamo 2008, Britanova 2008)",
    "- 2 Failed refinement: MEF2 limits excitatory synapse number and drives activity-dependent elimination (Flavell 2006; Pfeiffer 2010; Tsai 2012, via Pcdh10), so transient branches persist and grow from P5",
    "- 3 Receptor switch: Mef2c changes guidance receptors, e.g. Plxnd1 + Nrp1 turns Sema3e from repulsive to attractive (Chauvet 2007)",
    "- 4 Route: axons that fail to enter contralateral S1/S2 run on along the external capsule, which borders LA/BLA"], size=13.5, bold_first=True)
text(s, 6.9, 1.4, 6.0, 5.8, [
    "Target side / labelling",
    "- 5 Target change: an Emx1-type Cre also deletes Mef2c in BLA glutamatergic neurons (>90% Mef2c+ in several groups here)",
    "- 6 Compensation: losing the normal target frees axons and trophic competition, so they branch in an open neighbour",
    "Tests that separate them",
    "- Cre line; ipsi vs contra; time course of the route (P1-P5)",
    "- Retrograde CTB from BLA in the cKO -> source area and layer",
    "- Identity markers (Satb2, Cux2, Rorb, Bcl11b, Lmo4) in labelled cKO neurons",
    "- HCR at P5-P10 for the candidate cues (next slides)"], size=13.5, bold_first=True)

# 12b scenario 2
s = slide("Scenario 2 - BLA at P10 is a pre-pruning, pro-adhesion window", "Glial elimination machinery is still low at P10; target neurons express a P10-high adhesion programme")
pic(s, "F10_scenario2_pruning_window.png", 0.2, 1.3, 9.2, 6.1)
text(s, 9.5, 1.4, 3.6, 6.0, [
    "What the data show",
    "- Astrocyte Mertk 30% (P10) vs 75-77% (P21); Megf10 15% vs 26-33%",
    "- Microglial C1qa / C1qb 8% vs 17-37%; Trem2 10% vs 15-24%",
    "- Neurons: Ncam1, Sema6d, Cadm2, Pcdh7 higher at P10; Homer1, Nr4a1 (MEF2 targets) higher at P21 in Rspo2 BLA",
    "Reading",
    "- Ectopic axons arriving at P5-P10 meet a target that is not yet pruning",
    "- If Mef2c-null axons also lack their own MEF2-driven elimination, nothing removes them",
    "Caveat: one P10 animal; glial capture rises with age"], size=11.5, bold_first=True)

# 12c scenario 3 BLA side
s = slide("Scenario 3 (BLA side) - which cues peak in the P5-P10 window?", "* = same direction against both P21 samples")
pic(s, "F11_scenario3_BLA_cues.png", 0.2, 1.3, 7.9, 6.1)
text(s, 8.2, 1.4, 4.9, 6.0, [
    "Higher at P10",
    "- Sema6d: BLA Tshz2/Satb1, Rspo2/Etv1, Cdh8/Tshz3, ITC (a)",
    "- Slit2: Fgf10/Nrp1, Vgll3/Prr16, Rspo2/Etv1",
    "- Nrp1 in several BLA groups; Cdh8 in Rspo2/Etv1",
    "- Kirrel3 in ITC (b) (already 98%)",
    "Higher at P21 (maturing synapses)",
    "- Cdh9, Cdh18, C1ql3, Sdk2",
    "Sema3e is made by Rspo2 BLA at both ages (38% at P10); its receptor Plxnd1 is barely expressed inside the amygdala, so Sema3e acts on incoming axons"], size=12, bold_first=True)

# 12d scenario 3 cortex side
s = slide("Scenario 3 (cortex side) - partner receptors by cortical area", "Adult Allen WMB 10Xv2 (408,639 L2/3 + L4/5 IT cells). Numbers = % cells. * = area with known BLA input")
pic(s, "F12_scenario3_cortex_by_area.png", 0.2, 1.35, 12.9, 4.7)
text(s, 0.5, 6.1, 12.3, 1.3, [
    "- Cdh9 (homophilic; marks Rspo2 BLA): L2/3 IT 61-64% in TEa-PERI-ECT and PL-ILA-ORB, 46% in AI, but only 9% in SSp",
    "- Sema3e response: SSp L4/5 IT is mostly Plxnd1+ Nrp1- (35%, repulsion-type), with only 1.6% Plxnd1+ Nrp1+ (attraction-type)",
    "- Inside IT neurons Mef2c co-varies weakly and negatively with Nrp1, Cdh9 and Cdh18 (r = -0.05 to -0.08). Hypothesis: Mef2c loss de-represses Nrp1 / Cdh9 in SSp neurons, switching them toward BLA-type partners"], size=12)

# 12 summary
s = slide("Summary and suggested next steps")
text(s, 0.5, 1.2, 6.2, 6.1, [
    "Cell populations to watch",
    "- ITCs (Foxp2+ Tshz1+ Pbx3+), two groups: the FoxP2 cells next to the ectopic axons",
    "- Rspo2/Etv1 BLA glutamatergic neurons, the most distinct BLA group",
    "- Mef2c-high BLA groups (Tshz2/Satb1, Cdh8/Tshz3, Otof/Trhr): if the Cre also hits BLA neurons (e.g. Emx1 lineage), the target side changes too",
    "Gene candidates",
    "- ITC: Kirrel3, Sdk2, Sema6d, Sema5b, Cdh18, Ntng1",
    "- BLA Rspo2: C1ql3, Cdh8, Cdh9, Sema3e, Slit2",
    "- Cortex (test in cKO SSp L2/3-L4/5 at P5-P10): Nrp1 / Plxnd1 ratio, Cdh9, Kirrel3"], size=13, bold_first=True)
text(s, 6.9, 1.2, 6.0, 6.1, [
    "Next steps",
    "- Confirm which Cre line drives the cKO: this decides whether BLA neurons also lose Mef2c",
    "- HCR / RNAscope at P5-P10: Foxp2 + Kirrel3 / Sdk2 / Sema6d, and Rspo2 + Sema3e / Cdh9, next to the ectopic axons",
    "- Cortex: check the candidate partners (Kirrel3, Plxnd1, Cdh8/9, Sdk2) in WT vs Mef2c-cKO L2/3 neurons at P5-P10",
    "- Retrograde tracing from BLA in the cKO to find the source area, then profile it",
    "- New P10 replicates (P10_1, P10_2, run P173) are being processed and will make P10 vs P21 testable"], size=14, bold_first=True)

out = ROOT / "BLA_P10_P21_celltypes_TF_cues_for_Mef2c_v2.pptx"
prs.save(out); print("saved", out, len(prs.slides), "slides")
